from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook

from medreg.modules.documents.security import (
    EICAR_MARKER,
    ControlledFileSecurityInspector,
    FileSecurityError,
)

inspector = ControlledFileSecurityInspector()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("医疗器械注册资料")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    workbook = Workbook()
    workbook.active.append(["资料类别", "状态"])
    workbook.active.append(["风险分析", "已接受"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_mismatched_pdf_signature_is_blocked() -> None:
    with pytest.raises(FileSecurityError, match="does not match"):
        inspector.inspect("spoofed.pdf", "application/pdf", b"not a pdf")


def test_eicar_signature_is_blocked_before_storage() -> None:
    with pytest.raises(FileSecurityError, match="antivirus test signature"):
        inspector.inspect("sample.txt", "text/plain", EICAR_MARKER)


def test_safe_docx_and_xlsx_packages_receive_a_traceable_report() -> None:
    docx_report = inspector.inspect(
        "requirements.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        make_docx(),
    )
    xlsx_report = inspector.inspect(
        "matrix.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        make_xlsx(),
    )

    assert docx_report.status == "passed"
    assert xlsx_report.detected_type == "xlsx"
    assert "macro_free" in xlsx_report.findings


def test_macro_enabled_ooxml_package_is_blocked() -> None:
    source = make_docx()
    source_archive = ZipFile(BytesIO(source))
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source_archive.infolist():
            target.writestr(entry, source_archive.read(entry.filename))
        target.writestr("word/vbaProject.bin", b"macro payload")
    source_archive.close()

    with pytest.raises(FileSecurityError, match="Macro-enabled"):
        inspector.inspect("macro.docx", None, output.getvalue())
