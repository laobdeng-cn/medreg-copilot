from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import delete

from medreg.core.config import get_settings
from medreg.core.database import async_session_factory
from medreg.modules.documents.dispatcher import InMemoryDocumentTaskDispatcher
from medreg.modules.documents.fetcher import InMemoryOfficialSourceFetcher
from medreg.modules.documents.parser import ControlledDocumentParser
from medreg.modules.documents.repository import SQLAlchemyDocumentRepository
from medreg.modules.documents.schemas import ParseStatus
from medreg.modules.documents.segmenter import LegalDocumentSegmenter
from medreg.modules.documents.service import DocumentService
from medreg.modules.documents.storage import MinioObjectStorage
from medreg.modules.regulations.models import RegulationSourceModel
from medreg.modules.regulations.repository import SQLAlchemyRegulationRepository
from medreg.modules.regulations.schemas import (
    RegulationSourceCreate,
    RegulationType,
    RegulationVersionCreate,
)
from medreg.modules.regulations.service import RegulationService


@pytest.mark.integration
async def test_document_survives_database_and_minio_round_trip() -> None:
    settings = get_settings()
    storage = MinioObjectStorage(
        endpoint=settings.minio_endpoint,
        access_key=settings.minio_access_key,
        secret_key=settings.minio_secret_key,
        bucket_name=settings.minio_bucket,
        secure=settings.minio_secure,
    )
    source_id = None
    object_keys: list[str] = []
    payload = "第一章 总则\n第一条\n本办法适用于医疗器械注册与备案。".encode()

    try:
        async with async_session_factory() as write_session:
            regulation_service = RegulationService(
                SQLAlchemyRegulationRepository(write_session)
            )
            source = await regulation_service.create(
                RegulationSourceCreate(
                    title="文档归档集成测试法规",
                    issuing_authority="集成测试机关",
                    regulation_type=RegulationType.REGULATION,
                    scope_summary="验证 PostgreSQL 和 MinIO 的受控原文归档。",
                    initial_version=RegulationVersionCreate(
                        version_label="测试版",
                        document_number="TEST-DOC-001",
                        official_url="https://example.gov.cn/document-test",
                        published_on=date(2026, 1, 1),
                        effective_on=date(2026, 2, 1),
                    ),
                )
            )
            source_id = source.id

            document_service = DocumentService(
                repository=SQLAlchemyDocumentRepository(write_session),
                storage=storage,
                parser=ControlledDocumentParser(),
                segmenter=LegalDocumentSegmenter(),
                dispatcher=InMemoryDocumentTaskDispatcher(),
                fetcher=InMemoryOfficialSourceFetcher(),
                max_upload_bytes=settings.document_max_upload_bytes,
                parse_stale_after_seconds=settings.document_parse_stale_after_seconds,
            )
            archived = await document_service.archive(
                version_id=source.versions[0].id,
                file_name="integration-regulation.txt",
                content_type="text/plain",
                data=payload,
                uploaded_by="集成测试人员",
            )
            object_keys.append(archived.object_key)
            queued = await document_service.queue_parse(archived.id)
            parsed = await document_service.execute_parse(
                archived.id, queued.parse_task_id
            )
            assert parsed.parse_status.value == "completed"
            assert parsed.section_count == 2
            assert parsed.chunk_count == 1

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "资料矩阵"
            worksheet.append(["资料类别", "状态", "责任角色"])
            worksheet.append(["产品风险分析", "已接受", "法规专员"])
            worksheet.append(["临床评价资料", "待补正", "临床专员"])
            buffer = BytesIO()
            workbook.save(buffer)
            workbook.close()
            spreadsheet = await document_service.archive(
                version_id=source.versions[0].id,
                file_name="integration-dossier.xlsx",
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                data=buffer.getvalue(),
                uploaded_by="集成测试人员",
            )
            object_keys.append(spreadsheet.object_key)
            spreadsheet_queued = await document_service.queue_parse(spreadsheet.id)
            spreadsheet_parsed = await document_service.execute_parse(
                spreadsheet.id, spreadsheet_queued.parse_task_id
            )
            assert spreadsheet_parsed.parse_status is ParseStatus.COMPLETED
            assert spreadsheet_parsed.security_status.value == "passed"
            assert spreadsheet_parsed.detected_type == "xlsx"
            assert spreadsheet_parsed.table_count == 1

        async with async_session_factory() as read_session:
            repository = SQLAlchemyDocumentRepository(read_session)
            restored = await repository.get(archived.id)
            restored_structure = await repository.get_structure(archived.id)
            restored_spreadsheet = await repository.get(spreadsheet.id)
            restored_spreadsheet_structure = await repository.get_structure(
                spreadsheet.id
            )
            stale_write = await repository.update_parse(
                document_id=archived.id,
                status=ParseStatus.FAILED,
                attempts=99,
                task_id="superseded-task",
                parser_version="stale-worker",
                extracted_text=None,
                error="must not be written",
                queued_at=None,
                processing_started_at=None,
                parsed_at=None,
                updated_at=datetime.now(UTC),
                expected_task_id="superseded-task",
            )

        assert restored is not None
        assert restored_structure is not None
        assert restored_spreadsheet is not None
        assert restored_spreadsheet_structure is not None
        assert restored_structure.section_count == 2
        assert restored_structure.chunk_count == 1
        assert restored_structure.sections[1].parent_id == (
            restored_structure.sections[0].id
        )
        assert restored_structure.chunks[0].citation_label == "第一章 总则 / 第一条"
        assert stale_write is not None
        assert stale_write.parse_status == ParseStatus.COMPLETED
        assert stale_write.parse_attempts == 1
        assert restored.sha256 == archived.sha256
        assert restored.extracted_char_count > 10
        assert restored.parse_attempts == 1
        assert await storage.get(archived.object_key) == payload
        assert restored_spreadsheet.security_status.value == "passed"
        assert restored_spreadsheet.table_count == 1
        assert restored_spreadsheet_structure.table_count == 1
        assert restored_spreadsheet_structure.tables[0].title == "资料矩阵"
        assert restored_spreadsheet_structure.tables[0].headers == [
            "资料类别",
            "状态",
            "责任角色",
        ]
        assert restored_spreadsheet_structure.tables[0].rows[1] == [
            "临床评价资料",
            "待补正",
            "临床专员",
        ]
    finally:
        if source_id is not None:
            async with async_session_factory() as cleanup_session:
                await cleanup_session.execute(
                    delete(RegulationSourceModel).where(
                        RegulationSourceModel.id == source_id
                    )
                )
                await cleanup_session.commit()
        for object_key in object_keys:
            await storage.delete(object_key)
